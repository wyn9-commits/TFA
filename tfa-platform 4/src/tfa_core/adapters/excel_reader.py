"""tfa_core.adapters.excel_reader — Excel folios → DocumentLayout.

The ServiceNow data-type list includes Excel sheets. Spreadsheets don't go
through Document Intelligence; the workbook IS already structured, so we
render every sheet to a markdown table and feed text-only input to the same
Stage-2 LLM extraction (schema, validators, and repair loop unchanged).
"""
from __future__ import annotations

import io
import logging

from openpyxl import load_workbook

from tfa_core.domain.values import DocumentLayout

logger = logging.getLogger(__name__)

_MAX_ROWS_PER_SHEET = 500
_MAX_COLS = 40


def analyze_excel(content: bytes, filename: str) -> DocumentLayout:
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if r_idx >= _MAX_ROWS_PER_SHEET:
                rows.append(["… (sheet truncated)"])
                break
            cells = ["" if v is None else str(v) for v in row[:_MAX_COLS]]
            if any(c.strip() for c in cells):
                rows.append(cells)
        if not rows:
            continue
        width = max(len(r) for r in rows)
        norm = [r + [""] * (width - len(r)) for r in rows]
        md = "\n".join("| " + " | ".join(r) + " |" for r in norm)
        header_sep = "| " + " | ".join(["---"] * width) + " |"
        table = norm and (
            "| " + " | ".join(norm[0]) + " |\n" + header_sep + "\n"
            + "\n".join("| " + " | ".join(r) + " |" for r in norm[1:])
        ) or md
        parts.append(f"## Sheet: {ws.title}\n\n{table}")
    wb.close()

    markdown = "\n\n".join(parts) or "(workbook contains no data)"
    logger.info("Excel %s: %d sheet(s), %d chars markdown.",
                filename, len(parts), len(markdown))
    return DocumentLayout(
        markdown_text=markdown,
        page_count=len(parts),
        page_images_b64=[],          # text-only input for the LLM
        languages=[],
        ocr_confidence_floor=1.0,    # native data, no OCR uncertainty
    )
